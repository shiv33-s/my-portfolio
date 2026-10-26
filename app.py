
import re
from io import BytesIO
import fitz
from docx import Document
from flask import request, render_template

from flask import Flask, render_template, request, redirect, url_for, session, flash
from werkzeug.security import generate_password_hash, check_password_hash
import sqlite3, os
from functools import wraps

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-this-secret-key-for-production")
DB_PATH = os.path.join(os.path.dirname(__file__), "capacity_connect.db")

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def init_db():
    conn = db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS users(
      id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, email TEXT UNIQUE NOT NULL,
      password TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('trainee','trainer','admin')),
      skills TEXT DEFAULT '', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS courses(
      id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, description TEXT NOT NULL,
      category TEXT NOT NULL, skills TEXT NOT NULL, level TEXT DEFAULT 'Beginner',
      duration TEXT DEFAULT '2 hours', trainer_id INTEGER, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY(trainer_id) REFERENCES users(id)
    );
    CREATE TABLE IF NOT EXISTS enrollments(
      id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, course_id INTEGER NOT NULL,
      progress INTEGER DEFAULT 0, enrolled_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
      UNIQUE(user_id, course_id), FOREIGN KEY(user_id) REFERENCES users(id), FOREIGN KEY(course_id) REFERENCES courses(id)
    );
    CREATE TABLE IF NOT EXISTS questions(
      id INTEGER PRIMARY KEY AUTOINCREMENT, course_id INTEGER NOT NULL, question TEXT NOT NULL,
      option_a TEXT NOT NULL, option_b TEXT NOT NULL, option_c TEXT NOT NULL, option_d TEXT NOT NULL,
      answer TEXT NOT NULL CHECK(answer IN ('a','b','c','d')),
      FOREIGN KEY(course_id) REFERENCES courses(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS quiz_results(
      id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, course_id INTEGER NOT NULL,
      score INTEGER NOT NULL, total INTEGER NOT NULL, taken_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY(user_id) REFERENCES users(id), FOREIGN KEY(course_id) REFERENCES courses(id)
    );
    """)
    if conn.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
        conn.execute("INSERT INTO users(name,email,password,role,skills) VALUES(?,?,?,?,?)",
                     ("Portal Admin","admin@capacityconnect.local",generate_password_hash("Admin123!"),"admin","management,analytics"))
        trainer_id = conn.execute("INSERT INTO users(name,email,password,role,skills) VALUES(?,?,?,?,?)",
                     ("Demo Trainer","trainer@capacityconnect.local",generate_password_hash("Trainer123!"),"trainer","python,data analysis,communication")).lastrowid
        courses = [
          ("Introduction to Data Analysis","Learn how to interpret data, identify patterns and make evidence-based decisions.","Data","data analysis,excel,statistics","Beginner","3 hours"),
          ("Python Programming Basics","Understand variables, conditions, loops and functions through practical examples.","Technology","python,programming,problem solving","Beginner","4 hours"),
          ("Effective Workplace Communication","Build clear written, verbal and collaborative communication skills.","Professional Skills","communication,teamwork,leadership","Beginner","2 hours"),
          ("Foundations of Artificial Intelligence","Explore AI concepts, common applications and responsible use.","Technology","artificial intelligence,machine learning,data","Intermediate","3 hours")
        ]
        for c in courses:
            cid = conn.execute("INSERT INTO courses(title,description,category,skills,level,duration,trainer_id) VALUES(?,?,?,?,?,?,?)",
                               (c[0],c[1],c[2],c[3],c[4],c[5],trainer_id)).lastrowid
            if c[0] == "Introduction to Data Analysis":
                qs = [
                 ("What is the main purpose of data analysis?","To collect random facts","To inspect data and draw useful conclusions","To delete all records","To design logos","b"),
                 ("Which tool is commonly used for spreadsheet analysis?","Excel","Paint","Notepad only","Calculator only","a"),
                 ("What does a chart help communicate?","A visual pattern in data","A password","A file path","A programming error","a")]
                for q in qs:
                    conn.execute("INSERT INTO questions(course_id,question,option_a,option_b,option_c,option_d,answer) VALUES(?,?,?,?,?,?,?)",(cid,*q))
    conn.commit()
    conn.close()

init_db()

def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user_id"):
            flash("Please log in to continue.", "warning")
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapper

def role_required(*roles):
    def deco(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            if not session.get("user_id"):
                flash("Please log in to continue.", "warning")
                return redirect(url_for("login"))
            if session.get("role") not in roles:
                flash("You do not have permission to access that page.", "danger")
                return redirect(url_for("dashboard"))
            return fn(*args, **kwargs)
        return wrapper
    return deco

def current_user():
    if not session.get("user_id"): return None
    conn = db()
    user = conn.execute("SELECT * FROM users WHERE id=?", (session["user_id"],)).fetchone()
    conn.close()
    return user

@app.context_processor
def inject_user():
    return {"current_user": current_user()}

@app.route("/")
def home():
    conn = db()
    courses = conn.execute("SELECT courses.*, users.name AS trainer_name FROM courses LEFT JOIN users ON users.id=courses.trainer_id ORDER BY courses.id DESC LIMIT 3").fetchall()
    total_courses = conn.execute("SELECT COUNT(*) FROM courses").fetchone()[0]
    total_users = conn.execute("SELECT COUNT(*) FROM users WHERE role='trainee'").fetchone()[0]
    conn.close()
    return render_template("home.html", courses=courses, total_courses=total_courses, total_users=total_users)

@app.route("/register", methods=["GET","POST"])
def register():
    if request.method == "POST":
        name = request.form.get("name","").strip()
        email = request.form.get("email","").strip().lower()
        password = request.form.get("password","")
        skills = request.form.get("skills","").strip()
        if not name or not email or len(password) < 8:
            flash("Enter your name and email, and use a password of at least 8 characters.", "danger")
            return render_template("register.html")
        conn = db()
        try:
            conn.execute("INSERT INTO users(name,email,password,role,skills) VALUES(?,?,?,?,?)",
                         (name,email,generate_password_hash(password),"trainee",skills))
            conn.commit()
        except sqlite3.IntegrityError:
            conn.close()
            flash("An account with that email already exists.", "danger")
            return render_template("register.html")
        conn.close()
        flash("Account created! Log in to get started.", "success")
        return redirect(url_for("login"))
    return render_template("register.html")

@app.route("/login", methods=["GET","POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email","").strip().lower()
        password = request.form.get("password","")
        conn = db()
        user = conn.execute("SELECT * FROM users WHERE email=?", (email,)).fetchone()
        conn.close()
        if user and check_password_hash(user["password"], password):
            session.clear()
            session.update(user_id=user["id"], name=user["name"], role=user["role"])
            flash("Welcome back, " + user["name"] + "!", "success")
            return redirect(url_for("dashboard"))
        flash("Incorrect email or password.", "danger")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    flash("You have been logged out.", "info")
    return redirect(url_for("home"))

@app.route("/dashboard")
@login_required
def dashboard():
    conn = db()
    user = conn.execute("SELECT * FROM users WHERE id=?", (session["user_id"],)).fetchone()
    if user["role"] == "trainee":
        enrolled = conn.execute("""SELECT courses.*, enrollments.progress, users.name AS trainer_name
          FROM enrollments JOIN courses ON courses.id=enrollments.course_id
          LEFT JOIN users ON users.id=courses.trainer_id WHERE enrollments.user_id=? ORDER BY enrollments.id DESC""",(user["id"],)).fetchall()
        courses = conn.execute("SELECT * FROM courses ORDER BY title").fetchall()
        skill_tokens = {x.strip().lower() for x in (user["skills"] or "").split(",") if x.strip()}
        recs = []
        for c in courses:
            tags = {x.strip().lower() for x in c["skills"].split(",") if x.strip()}
            overlap = skill_tokens & tags
            # A low overlap indicates a potential learning opportunity; exact match is a simple prototype recommendation.
            if c["id"] not in [e["id"] for e in enrolled] and (not skill_tokens or overlap or len(recs)<2):
                recs.append(dict(course=c, match=len(overlap), missing=sorted(tags-skill_tokens)))
        recs.sort(key=lambda x: (x["match"], x["course"]["title"]))
        recs = recs[:3]
        results = conn.execute("""SELECT quiz_results.*, courses.title FROM quiz_results JOIN courses ON courses.id=quiz_results.course_id
          WHERE quiz_results.user_id=? ORDER BY quiz_results.id DESC LIMIT 5""",(user["id"],)).fetchall()
        conn.close()
        return render_template("trainee_dashboard.html", user=user, enrolled=enrolled, recommendations=recs, results=results)
    elif user["role"] == "trainer":
        courses = conn.execute("SELECT * FROM courses WHERE trainer_id=? ORDER BY id DESC",(user["id"],)).fetchall()
        stats = conn.execute("""SELECT COUNT(*) FROM enrollments JOIN courses ON courses.id=enrollments.course_id WHERE courses.trainer_id=?""",(user["id"],)).fetchone()[0]
        conn.close()
        return render_template("trainer_dashboard.html", user=user, courses=courses, enrollments_count=stats)
    else:
        users = conn.execute("SELECT id,name,email,role,skills,created_at FROM users ORDER BY id DESC").fetchall()
        courses = conn.execute("SELECT courses.*, users.name AS trainer_name FROM courses LEFT JOIN users ON users.id=courses.trainer_id ORDER BY courses.id DESC").fetchall()
        enrollments = conn.execute("SELECT COUNT(*) FROM enrollments").fetchone()[0]
        conn.close()
        return render_template("admin_dashboard.html", user=user, users=users, courses=courses, enrollments=enrollments)

@app.route("/courses")
def course_list():
    q = request.args.get("q","").strip()
    conn = db()
    courses = conn.execute("""SELECT courses.*, users.name AS trainer_name FROM courses LEFT JOIN users ON users.id=courses.trainer_id
       WHERE courses.title LIKE ? OR courses.description LIKE ? OR courses.category LIKE ? ORDER BY courses.title""",
       (f"%{q}%",f"%{q}%",f"%{q}%")).fetchall()
    enrolled_ids = []
    if session.get("user_id"):
        enrolled_ids = [r["course_id"] for r in conn.execute("SELECT course_id FROM enrollments WHERE user_id=?",(session["user_id"],)).fetchall()]
    conn.close()
    return render_template("courses.html", courses=courses, enrolled_ids=enrolled_ids, q=q)

@app.route("/courses/<int:course_id>")
def course_detail(course_id):
    conn = db()
    course = conn.execute("SELECT courses.*, users.name AS trainer_name FROM courses LEFT JOIN users ON users.id=courses.trainer_id WHERE courses.id=?",(course_id,)).fetchone()
    questions = conn.execute("SELECT id,question,option_a,option_b,option_c,option_d FROM questions WHERE course_id=?",(course_id,)).fetchall()
    is_enrolled = False
    if session.get("user_id"):
        is_enrolled = conn.execute("SELECT 1 FROM enrollments WHERE user_id=? AND course_id=?",(session["user_id"],course_id)).fetchone() is not None
    conn.close()
    if not course:
        flash("Course not found.", "danger")
        return redirect(url_for("course_list"))
    return render_template("course_detail.html", course=course, questions=questions, is_enrolled=is_enrolled)

@app.route("/courses/<int:course_id>/enroll", methods=["POST"])
@role_required("trainee")
def enroll(course_id):
    conn = db()
    try:
        conn.execute("INSERT INTO enrollments(user_id,course_id) VALUES(?,?)",(session["user_id"],course_id))
        conn.commit()
        flash("You are enrolled in this course.", "success")
    except sqlite3.IntegrityError:
        flash("You are already enrolled in this course.", "info")
    conn.close()
    return redirect(url_for("course_detail",course_id=course_id))

@app.route("/courses/<int:course_id>/quiz", methods=["POST"])
@role_required("trainee")
def take_quiz(course_id):
    conn = db()
    enrolled = conn.execute("SELECT 1 FROM enrollments WHERE user_id=? AND course_id=?",(session["user_id"],course_id)).fetchone()
    questions = conn.execute("SELECT * FROM questions WHERE course_id=?",(course_id,)).fetchall()
    if not enrolled:
        conn.close()
        flash("Enroll in the course before attempting its quiz.", "warning")
        return redirect(url_for("course_detail",course_id=course_id))
    if not questions:
        conn.close()
        flash("There is no quiz available for this course yet.", "info")
        return redirect(url_for("course_detail",course_id=course_id))
    score = sum(1 for q in questions if request.form.get(str(q["id"])) == q["answer"])
    total = len(questions)
    conn.execute("INSERT INTO quiz_results(user_id,course_id,score,total) VALUES(?,?,?,?)",(session["user_id"],course_id,score,total))
    conn.execute("UPDATE enrollments SET progress=MAX(progress,100) WHERE user_id=? AND course_id=?",(session["user_id"],course_id))
    conn.commit()
    conn.close()
    flash(f"Quiz submitted. Your score is {score}/{total}.", "success")
    return redirect(url_for("dashboard"))

@app.route("/trainer/courses/new", methods=["GET","POST"])
@role_required("trainer")
def new_course():
    if request.method == "POST":
        title=request.form.get("title","").strip()
        description=request.form.get("description","").strip()
        category=request.form.get("category","").strip()
        skills=request.form.get("skills","").strip()
        level=request.form.get("level","Beginner")
        duration=request.form.get("duration","2 hours").strip()
        if not all([title,description,category,skills]):
            flash("Please complete all required fields.", "danger")
            return render_template("new_course.html")
        conn=db()
        conn.execute("INSERT INTO courses(title,description,category,skills,level,duration,trainer_id) VALUES(?,?,?,?,?,?,?)",
                     (title,description,category,skills,level,duration,session["user_id"]))
        conn.commit(); conn.close()
        flash("Course created successfully.", "success")
        return redirect(url_for("dashboard"))
    return render_template("new_course.html")

@app.route("/admin/users/<int:user_id>/role", methods=["POST"])
@role_required("admin")
def update_role(user_id):
    role=request.form.get("role")
    if role not in ("trainee","trainer","admin"):
        flash("Invalid role.", "danger")
    elif user_id == session["user_id"]:
        flash("You cannot change your own role.", "warning")
    else:
        conn=db()
        conn.execute("UPDATE users SET role=? WHERE id=?",(role,user_id))
        conn.commit(); conn.close()
        flash("User role updated.", "success")
    return redirect(url_for("dashboard"))

@app.errorhandler(404)
def not_found(error):
    return render_template("404.html"), 404

# ---------- RESUME ANALYZER ----------

ALLOWED_RESUME_EXTENSIONS = {"pdf", "docx"}
MAX_RESUME_SIZE = 5 * 1024 * 1024

SKILL_KEYWORDS = [
    "python", "java", "c++", "sql", "mysql", "sqlite",
    "html", "css", "javascript", "flask", "django",
    "machine learning", "deep learning",
    "artificial intelligence", "data analysis",
    "data visualization", "pandas", "numpy",
    "scikit-learn", "tensorflow", "pytorch", "nlp",
    "computer vision", "excel", "power bi", "tableau",
    "git", "github", "communication", "problem solving",
    "teamwork", "leadership", "statistics",
    "cloud computing", "rest api", "mongodb",
    "react", "node.js", "c", "c++"
]


def extract_resume_text(file_bytes, extension):
    if extension == "pdf":
        with fitz.open(
            stream=file_bytes,
            filetype="pdf"
        ) as pdf:
            return "\n".join(
                page.get_text() for page in pdf
            )

    if extension == "docx":
        document = Document(BytesIO(file_bytes))

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)

        return "\n".join(parts)

    raise ValueError("Unsupported file format")


def extract_skills(text):
    text = text.lower()
    found = []

    for skill in SKILL_KEYWORDS:
        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(skill)
            + r"(?![a-z0-9])"
        )

        if re.search(pattern, text):
            found.append(skill)

    return sorted(set(found))


@app.route("/resume-analyzer", methods=["GET", "POST"])
def resume_analyzer():
    result = None
    error = None

    if request.method == "POST":
        uploaded_file = request.files.get("resume")
        job_description = request.form.get(
            "job_description", ""
        ).strip()

        if not uploaded_file or not uploaded_file.filename:
            error = "Please select a resume."

        elif "." not in uploaded_file.filename:
            error = "Please upload a PDF or DOCX file."

        else:
            extension = uploaded_file.filename.rsplit(
                ".", 1
            )[1].lower()

            if extension not in ALLOWED_RESUME_EXTENSIONS:
                error = "Only PDF and DOCX files are supported."

            elif not job_description:
                error = "Please enter the job description."

            else:
                file_bytes = uploaded_file.read(
                    MAX_RESUME_SIZE + 1
                )

                if len(file_bytes) > MAX_RESUME_SIZE:
                    error = "Resume size must be 5 MB or less."

                else:
                    try:
                        resume_text = extract_resume_text(
                            file_bytes, extension
                        )

                        if not resume_text.strip():
                            error = (
                                "No readable text found. "
                                "Try a text-based PDF or DOCX."
                            )
                        else:
                            resume_skills = extract_skills(
                                resume_text
                            )
                            job_skills = extract_skills(
                                job_description
                            )

                            matched = [
                                skill for skill in job_skills
                                if skill in resume_skills
                            ]

                            missing = [
                                skill for skill in job_skills
                                if skill not in resume_skills
                            ]

                            percentage = (
                                round(
                                    len(matched)
                                    / len(job_skills) * 100
                                )
                                if job_skills else 0
                            )

                            result = {
                                "resume_skills": resume_skills,
                                "job_skills": job_skills,
                                "matched_skills": matched,
                                "missing_skills": missing,
                                "match_percentage": percentage,
                                "word_count": len(
                                    resume_text.split()
                                ),
                                "job_skill_count": len(job_skills)
                            }

                    except Exception:
                        app.logger.exception(
                            "Resume analysis failed"
                        )
                        error = (
                            "Unable to process this document. "
                            "Please upload a valid PDF or DOCX."
                        )

    return render_template(
        "resume_analyzer.html",
        result=result,
        error=error
    )

if __name__ == "__main__":
    app.run(debug=True)
